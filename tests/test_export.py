"""Exporting order history to a spreadsheet.

The file is read outside this project, by someone deciding whether a strategy
works. So the tests are about it saying true things: the right day, the right
period, and no number that was estimated rather than measured.
"""

from __future__ import annotations

import io
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import pytest

#: The project root. These tests read source files as text, so the
#: path is resolved from this rather than repeated at each use.
PROJECT_ROOT = Path(__file__).resolve().parent.parent

sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402

from backend.api.errors import ApiError  # noqa: E402
from backend.api.routes.export import market_day_of, parse_day  # noqa: E402
from backend.services.export import build_order_row, build_workbook  # noqa: E402
from backend.services.export.columns import ORDER_COLUMNS  # noqa: E402


def an_order(**overrides) -> dict:
    """One order, shaped as GET /orders/history returns it."""
    order = {
        "order_id_text": "44674820673913856",
        "placed_at": datetime(2026, 9, 18, 17, 34, tzinfo=timezone.utc),
        "filled_at": datetime(2026, 9, 18, 17, 34, 1, tzinfo=timezone.utc),
        "exited_at": datetime(2026, 9, 18, 17, 36, tzinfo=timezone.utc),
        "identifier": "NVDA  260921C00215000",
        "underlying": "NVDA",
        "strike": 215.0,
        "option_type": "CALL",
        "expiry": "2026-09-21",
        "quantity": 2.0,
        "limit_price": 2.88,
        "fill_price": 2.85,
        "exit_price": 3.02,
        "status": "FILLED",
        "outcome": "TOOK_PROFIT",
        "outcome_note": "TAKE PROFIT triggered.",
        "realised_pnl": 34.0,
        "realised_pnl_percent": 5.96,
        "take_profit_price": 3.03,
        "stop_loss_price": 2.59,
        "leg_time_in_force": "DAY",
        "filled_at_seconds": None,
        "fill_delay_seconds": 1.0,
        "held_seconds": 119.0,
        "held_open_seconds": None,
        "legs": [],
    }
    order.update(overrides)
    return order


class TestTheDayIsTheMarketsDay:
    """A trading session belongs to New York, not to the reader.

    The history page's date filter had this bug: from UTC+5:30 every US
    afternoon trade landed on the NEXT calendar day, so picking today found
    nothing. The export groups and names files by day, so it would be the
    same bug in a file someone keeps.
    """

    def test_a_us_afternoon_trade_stays_on_its_own_day(self):
        """19:45 UTC is 15:45 in New York -- the same trading day, though it
        is already tomorrow in much of the world."""
        stamp = datetime(2026, 9, 18, 19, 45, tzinfo=timezone.utc)

        assert market_day_of(stamp) == date(2026, 9, 18)

    def test_an_iso_string_is_accepted_too(self):
        assert market_day_of("2026-09-18T19:45:00Z") == date(2026, 9, 18)

    def test_a_late_evening_utc_trade_is_the_same_us_day(self):
        """23:30 UTC is 19:30 in New York -- still the 18th there."""
        stamp = datetime(2026, 9, 18, 23, 30, tzinfo=timezone.utc)

        assert market_day_of(stamp) == date(2026, 9, 18)

    def test_nothing_returns_none(self):
        assert market_day_of(None) is None
        assert market_day_of("") is None

    def test_a_malformed_stamp_returns_none_rather_than_raising(self):
        """One bad row must not fail the whole export."""
        assert market_day_of("not a date") is None


class TestDateParsing:
    """A malformed date must not quietly export a period nobody asked for."""

    def test_a_valid_date_is_read(self):
        assert parse_day("2026-09-18", "start") == date(2026, 9, 18)

    def test_nothing_returns_none(self):
        assert parse_day(None, "start") is None
        assert parse_day("", "start") is None

    def test_the_wrong_order_is_refused(self):
        """18-09-2026 is a real date written the other way round. Guessing
        which it meant would export the wrong period."""
        with pytest.raises(ApiError) as caught:
            parse_day("18-09-2026", "start")

        assert caught.value.error_code == "DATE_INVALID"

    def test_an_impossible_date_is_refused(self):
        with pytest.raises(ApiError):
            parse_day("2026-02-30", "start")

    def test_the_message_names_the_parameter(self):
        with pytest.raises(ApiError) as caught:
            parse_day("nonsense", "end")

        assert "end" in caught.value.message


class TestTheRow:
    """Each row carries the numbers behind it, not just the conclusion."""

    def row(self, **overrides):
        return dict(zip(
            [heading for heading, _ in ORDER_COLUMNS],
            build_order_row(an_order(**overrides)),
        ))

    def test_the_date_is_the_market_day(self):
        assert self.row()["Date"] == date(2026, 9, 18)

    def test_slippage_is_the_fill_against_the_limit(self):
        """Negative means it filled BELOW the limit, which says the buffer was
        wider than it needed to be."""
        assert self.row()["Slippage"] == pytest.approx(-0.03)

    def test_cash_in_multiplies_by_the_contract_size(self):
        """2.85 x 2 contracts x 100 shares."""
        assert self.row()["Cash in"] == pytest.approx(570.0)

    def test_the_exit_distances_are_measured_from_the_fill(self):
        row = self.row()

        assert row["TP distance %"] == pytest.approx(6.32, abs=0.01)
        assert row["SL distance %"] == pytest.approx(-9.12, abs=0.01)

    def test_held_falls_back_to_the_open_clock(self):
        """An open position has no exit to measure to, so its running clock
        is the honest answer rather than a blank."""
        row = self.row(held_seconds=None, held_open_seconds=300.0)

        assert row["Held (s)"] == 300.0
        assert row["Held (min)"] == pytest.approx(5.0)

    def test_a_missing_price_does_not_break_the_row(self):
        """An order that never filled still belongs in the file."""
        row = self.row(fill_price=None, limit_price=None, realised_pnl=None)

        assert row["Slippage"] is None
        assert row["Cash in"] is None
        assert row["TP distance %"] is None

    def test_every_column_is_filled(self):
        """The row and the headings must not drift apart."""
        assert len(build_order_row(an_order())) == len(ORDER_COLUMNS)


class TestTheWorkbook:
    """It has to open, and say what period it covers."""

    def book(self, orders, start=date(2026, 9, 18), end=date(2026, 9, 18)):
        return openpyxl.load_workbook(
            io.BytesIO(build_workbook(orders, start, end))
        )

    def test_it_has_the_three_sheets(self):
        assert self.book([an_order()]).sheetnames == ["Orders", "Summary", "Legs"]

    def test_an_empty_period_still_produces_a_file(self):
        """A file that says "nothing happened" beats a failed download."""
        book = self.book([])

        assert book["Orders"].max_row == 1  # headings only

    def test_orders_are_written_oldest_first(self):
        """A spreadsheet is read downwards and a session runs forwards. The
        page shows newest first for the opposite reason."""
        older = an_order(placed_at=datetime(2026, 9, 18, 14, 0, tzinfo=timezone.utc),
                         underlying="AAPL")
        newer = an_order(placed_at=datetime(2026, 9, 18, 18, 0, tzinfo=timezone.utc),
                         underlying="TSLA")

        # The endpoint hands them over newest first.
        sheet = self.book([newer, older])["Orders"]

        assert sheet.cell(row=2, column=5).value == "AAPL"
        assert sheet.cell(row=3, column=5).value == "TSLA"

    def test_the_summary_names_the_period(self):
        summary = self.book([], date(2026, 9, 14), date(2026, 9, 18))["Summary"]

        assert summary.cell(row=1, column=2).value == "2026-09-14 to 2026-09-18"

    def test_a_single_day_summary_names_one_date(self):
        summary = self.book([])["Summary"]

        assert summary.cell(row=1, column=2).value == "2026-09-18"

    def test_the_legs_sheet_carries_one_row_per_leg(self):
        order = an_order(legs=[
            {"role": "TAKE_PROFIT", "order_type": "LMT", "price": 3.03,
             "time_in_force": "DAY", "status": "FILLED", "order_id_text": "1"},
            {"role": "STOP_LOSS", "order_type": "STP", "price": 2.59,
             "time_in_force": "DAY", "status": "CANCELLED", "order_id_text": "2"},
        ])

        assert self.book([order])["Legs"].max_row == 3  # headings + 2


class TestNothingIsEstimated:
    """Commission is not on the order record.

    A commission column computed from a rate would read as a measurement while
    being an assumption, and the file exists to decide whether a strategy
    works -- an invented cost would decide it wrongly.
    """

    def test_no_commission_column(self):
        headings = [heading for heading, _ in ORDER_COLUMNS]

        assert not any("commission" in h.lower() for h in headings)

    def test_the_summary_says_the_pnl_is_gross(self):
        summary = build_workbook([], date(2026, 9, 18), date(2026, 9, 18))
        book = openpyxl.load_workbook(io.BytesIO(summary))
        text = " ".join(
            str(cell.value or "")
            for row in book["Summary"].iter_rows()
            for cell in row
        )

        assert "gross" in text.lower()
        assert "commission" in text.lower()

    def test_the_export_package_imports_nothing_that_can_trade(self):
        """It is kept apart from the trading path on purpose."""
        for name in ("columns.py", "workbook.py", "__init__.py"):
            source = (PROJECT_ROOT / "backend/services/export" / name).read_text(
                encoding="utf-8"
            )
            assert "place_order" not in source
            assert "sell_option" not in source
            assert "trade_client" not in source


class TestRealisedFromFills:
    """The figure the broker itself computes.

    The order history pairs each buy with the sell it thinks closed it, and
    that pairing over-counts once a contract is traded more than once. Checked
    against the account's own trade-history export: NVDA 220 CALL traded six
    times made +46, and the paired history claimed +466 because two buys had
    been handed a later trade's exit price. The day was +451; the paired
    figure said +1153.
    """

    def fill(self, action, amount, quantity=1, contract="NVDA  260921C00220000",
             trade_time=1789740000000):
        from types import SimpleNamespace

        return SimpleNamespace(
            identifier=contract,
            action=action,
            filled_cash_amount=amount,
            filled=quantity,
            trade_time=trade_time,
            order_time=trade_time,
        )

    def totals(self, fills, day=date(2026, 9, 18)):
        from backend.services.export.realised import realised_by_contract

        return realised_by_contract(fills, day, day)

    def test_a_round_trip_is_the_difference(self):
        """Bought 546, sold 666 -- made 120, whatever the prices were."""
        totals = self.totals([
            self.fill("BUY", 546.0, 3),
            self.fill("SELL", 666.0, 3),
        ])

        assert totals["NVDA  260921C00220000"]["realised"] == 120.0

    def test_trading_the_same_contract_six_times_does_not_over_count(self):
        """The bug this module exists for. Six round trips netting +46."""
        fills = []
        for bought, sold in (
            (546.0, 483.0), (534.0, 522.0), (612.0, 612.0),
            (612.0, 652.0), (513.0, 546.0), (618.0, 666.0),
        ):
            fills.append(self.fill("BUY", bought, 3))
            fills.append(self.fill("SELL", sold, 3))

        assert self.totals(fills)["NVDA  260921C00220000"]["realised"] == 46.0

    def test_a_loss_is_negative(self):
        totals = self.totals([
            self.fill("BUY", 700.0, 2),
            self.fill("SELL", 600.0, 2),
        ])

        assert totals["NVDA  260921C00220000"]["realised"] == -100.0

    def test_contracts_are_kept_apart(self):
        totals = self.totals([
            self.fill("BUY", 100.0, 1, contract="AAA  1"),
            self.fill("SELL", 150.0, 1, contract="AAA  1"),
            self.fill("BUY", 200.0, 1, contract="BBB  2"),
            self.fill("SELL", 180.0, 1, contract="BBB  2"),
        ])

        assert totals["AAA  1"]["realised"] == 50.0
        assert totals["BBB  2"]["realised"] == -20.0

    def test_an_unclosed_position_realises_nothing(self):
        """Bought and not sold. Nothing has been realised -- the cost is
        reported separately rather than counted as a loss, which is what made
        the page's total swing negative whenever a position was open."""
        totals = self.totals([self.fill("BUY", 546.0, 3)])
        entry = totals["NVDA  260921C00220000"]

        assert entry["open_qty"] == 3
        assert entry["realised"] == 0.0
        assert entry["open_cost"] == 546.0

    def test_fills_outside_the_period_are_excluded(self):
        from datetime import datetime
        from zoneinfo import ZoneInfo

        ny = ZoneInfo("America/New_York")
        yesterday = int(
            datetime(2026, 9, 17, 12, tzinfo=ny).timestamp() * 1000
        )

        totals = self.totals([
            self.fill("BUY", 546.0, 3, trade_time=yesterday),
            self.fill("SELL", 666.0, 3),
        ])

        # Yesterday's buy is out of range, so the sell has nothing in the
        # period to close: it opens a short rather than realising anything.
        # A period boundary cannot invent a profit.
        assert totals["NVDA  260921C00220000"]["realised"] == 0.0
        assert totals["NVDA  260921C00220000"]["open_qty"] == -3

    def test_the_day_is_the_market_day(self):
        """A fill at 23:30 UTC is still the same US trading day."""
        from datetime import datetime, timezone

        from backend.services.export.realised import market_day_of

        late = self.fill(
            "SELL", 100.0,
            trade_time=int(
                datetime(2026, 9, 18, 23, 30, tzinfo=timezone.utc).timestamp() * 1000
            ),
        )

        assert market_day_of(late) == date(2026, 9, 18)

    def test_the_total_adds_the_contracts_up(self):
        from backend.services.export.realised import realised_total

        totals = self.totals([
            self.fill("BUY", 100.0, 1, contract="AAA  1"),
            self.fill("SELL", 150.0, 1, contract="AAA  1"),
            self.fill("BUY", 200.0, 1, contract="BBB  2"),
            self.fill("SELL", 260.0, 1, contract="BBB  2"),
        ])

        assert realised_total(totals) == 110.0

    def test_a_fill_with_no_quantity_is_ignored(self):
        """A cancelled order reported with zero filled must not move a total."""
        totals = self.totals([
            self.fill("BUY", 546.0, 3),
            self.fill("SELL", 0.0, 0),
        ])

        assert totals["NVDA  260921C00220000"]["fills"] == 1


class TestAnOpenPositionIsNotALoss:
    """Realised means CLOSED.

    Cash in minus cash out is right only when a position ended flat. While
    contracts are still held the cash has gone out and none has come back, so
    that formula reports the PURCHASE as a loss: a position bought for 584 and
    still held showed -584, and the page's total swung from +1153 to -129 as
    soon as anything was open.

    The fix matches sells to the buys they closed, first in first out -- how a
    broker accounts for it. Contracts still held simply do not contribute.
    """

    def fill(self, action, amount, quantity, contract="X  1"):
        from types import SimpleNamespace

        self._clock = getattr(self, "_clock", 0) + 1000
        stamp = 1789740000000 + self._clock
        return SimpleNamespace(
            identifier=contract,
            action=action,
            filled_cash_amount=amount,
            filled=quantity,
            trade_time=stamp,
            order_time=stamp,
        )

    def total(self, fills):
        from backend.services.export.realised import realised_by_contract, realised_total

        day = date(2026, 9, 18)
        return realised_total(realised_by_contract(fills, day, day))

    def test_an_open_position_contributes_nothing(self):
        """Bought and not yet sold. Nothing has been realised, so the total is
        zero -- not minus what it cost."""
        assert self.total([self.fill("BUY", 584.0, 2)]) == 0.0

    def test_a_closed_trade_beside_an_open_one_is_unaffected(self):
        """The bug that made the tile swing: opening a new position must not
        change what an already-finished trade made."""
        closed_only = self.total([
            self.fill("BUY", 584.0, 2),
            self.fill("SELL", 610.0, 2),
        ])
        with_an_open_one = self.total([
            self.fill("BUY", 584.0, 2),
            self.fill("SELL", 610.0, 2),
            self.fill("BUY", 600.0, 2),
        ])

        assert closed_only == 26.0
        assert with_an_open_one == 26.0

    def test_a_partial_close_realises_only_the_part_sold(self):
        """Bought 4 at 100 each, sold 2 at 110 each."""
        assert self.total([
            self.fill("BUY", 400.0, 4),
            self.fill("SELL", 220.0, 2),
        ]) == 20.0

    def test_sells_match_the_oldest_buys_first(self):
        """FIFO. Averaging every buy together instead would blend the cost of
        a position still held into a trade that already finished."""
        assert self.total([
            self.fill("BUY", 100.0, 1),   # cheap, bought first
            self.fill("BUY", 200.0, 1),   # dear, bought second
            self.fill("SELL", 150.0, 1),  # closes the CHEAP one: +50
        ]) == 50.0

    def test_a_short_is_not_realised_until_bought_back(self):
        """Sold more than held. The proceeds are not a profit yet."""
        assert self.total([self.fill("SELL", 500.0, 2)]) == 0.0

    def test_buying_back_a_short_realises_it(self):
        """Sold 2 for 500, bought them back for 450 -- made 50."""
        assert self.total([
            self.fill("SELL", 500.0, 2),
            self.fill("BUY", 450.0, 2),
        ]) == 50.0

    def test_the_open_cost_is_reported_separately(self):
        """So a reader can see what was excluded rather than wondering."""
        from backend.services.export.realised import realised_by_contract

        day = date(2026, 9, 18)
        totals = realised_by_contract([self.fill("BUY", 584.0, 2)], day, day)

        assert totals["X  1"]["open_cost"] == 584.0
        assert totals["X  1"]["open_qty"] == 2

    def test_contracts_do_not_borrow_each_others_lots(self):
        """A sell on one contract must not close a buy on another."""
        assert self.total([
            self.fill("BUY", 100.0, 1, contract="AAA  1"),
            self.fill("SELL", 150.0, 1, contract="BBB  2"),
        ]) == 0.0
